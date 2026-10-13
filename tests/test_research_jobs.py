from copy import deepcopy
import json
from pathlib import Path

import pandas as pd
import pytest

from core import research_jobs as jobs
from core import task_decomposition as task
from core.execution_budget import BudgetExceeded
from core.storage import configure_public_session, configure_workspace, get_data_root, resolve_stored_file, workspace_context


QUESTION = '筛选2024年营收增长但现金流下降的企业，并解释原因'


def frames(*names):
    return {name: pd.DataFrame([{'year': 2023, 'revenue': 100, 'operating_cashflow': 20},
                               {'year': 2024, 'revenue': 120, 'operating_cashflow': 15}]) for name in names}


@pytest.fixture
def library(tmp_path, monkeypatch):
    monkeypatch.setenv('FINANCIAL_DEPLOYMENT', 'local')
    monkeypatch.setenv('FINANCIAL_DATA_DIR', str(tmp_path))
    with workspace_context('evaluation', tmp_path / 'testing', allow_writes=True):
        get_data_root().mkdir()
        yield get_data_root()


def test_budget_boundary_resumes_only_unfinished_companies(library, monkeypatch):
    calls = []
    def retrieve(question, context, **kwargs):
        name = context['companies'][0]
        calls.append(name)
        return {'status': 'success', 'answer': name + '原文解释', 'citations': [{'company': name, 'snippet': '原文'}]}
    monkeypatch.setattr(task, 'retrieve_documents', retrieve)
    job = jobs.create_research_job(QUESTION, frames('甲', '乙'))
    def stop(stage):
        if stage.startswith('企业检索检查点已保存'):
            raise BudgetExceeded('测试步骤预算')
    paused = jobs.resume_research_job(job['job_id'], check_interrupt=stop)
    assert paused['status'] == 'paused' and paused['completed'] == 1 and paused['result'] is None
    assert calls == ['甲']
    done = jobs.resume_research_job(job['job_id'])
    assert done['status'] == 'completed' and calls == ['甲', '乙']
    assert [row['explanation'] for row in done['result']['rows']] == ['甲原文解释', '乙原文解释']
    assert done['result']['research_job_id'] == job['job_id']
    assert jobs.resume_research_job(job['job_id'])['result'] == done['result']
    assert calls == ['甲', '乙']


def test_freezes_values_and_csv_provenance_independent_of_new_library(library):
    source = library / 'data.csv'
    source.write_text('year,revenue\n2024,120\n', encoding='utf-8')
    data = frames('甲')
    data['甲']['metric_provenance'] = [{'revenue': {'file_path': str(source), 'cell': 'B2'}}] * 2
    reports = [{'id': 1, 'company_name': '甲', 'report_year': 2024, 'file_path': str(source), 'file_name': 'data.csv',
                'source_type': 'local_upload', 'source_url': 'https://example.org/report.csv', 'content_hash': 'original-digest'}]
    original = deepcopy(data['甲'].to_dict('records'))
    job = jobs.create_research_job(QUESTION, data, reports=reports)
    assert jobs.create_research_job(QUESTION, data, reports=reports)['job_id'] == job['job_id']
    source.write_text('year,revenue\n2024,999\n', encoding='utf-8')
    data['甲'].loc[1, 'revenue'] = 999
    assert jobs.research_input_fingerprint(data, reports) != job['source_fingerprint']
    result = jobs.resume_research_job(job['job_id'])['result']
    assert result['rows'][0]['revenue_current'] == 120
    assert result['report_files_snapshot'][0]['source_type'] == 'local_upload'
    assert result['report_files_snapshot'][0]['source_url'] == 'https://example.org/report.csv'
    assert result['report_files_snapshot'][0]['content_hash'] == 'original-digest'
    saved = Path(result['data_snapshot']['甲'][1]['metric_provenance']['revenue']['file_path'])
    assert saved.read_text() == 'year,revenue\n2024,120\n'
    assert saved != source and resolve_stored_file(str(saved)) == saved
    assert original[1]['revenue'] == 120


def test_real_docx_retrieval_after_original_is_removed(library):
    from docx import Document
    company = '测试动力股份有限公司'
    path = library / 'annual.docx'
    document = Document()
    document.add_paragraph(company + ' 2024年年度报告')
    document.add_paragraph('经营现金流下降主要由于本期采购款增加。')
    document.save(path)
    reports = [{'id': 1, 'company_name': company, 'report_year': 2024, 'file_type': 'docx',
                'file_path': str(path), 'file_name': 'annual.docx'}]
    job = jobs.create_research_job(QUESTION, frames(company), reports=reports)
    path.unlink()
    result = jobs.resume_research_job(job['job_id'])['result']
    assert result['retrieval_result']['citations']
    citation = result['retrieval_result']['citations'][0]
    assert '采购款增加' in citation['snippet']
    assert resolve_stored_file(citation['file_path']).is_file()


def test_exceptions_preserve_last_real_checkpoint(library, monkeypatch):
    calls = []
    def retrieve(question, context, **kwargs):
        name = context['companies'][0]
        calls.append(name)
        if name == '乙' and calls.count(name) == 1:
            raise RuntimeError('provider error with potentially private text')
        return {'answer': name, 'status': 'success'}
    monkeypatch.setattr(task, 'retrieve_documents', retrieve)
    job = jobs.create_research_job(QUESTION, frames('甲', '乙'))
    paused = jobs.resume_research_job(job['job_id'])
    assert paused['completed'] == 1 and paused['error'] == 'RuntimeError'
    assert 'private' not in (library / 'research_jobs' / job['job_id'] / 'job.json').read_text(encoding='utf-8')
    assert jobs.resume_research_job(job['job_id'])['status'] == 'completed'
    assert calls == ['甲', '乙', '乙']


def test_process_interruption_records_pause_and_releases_lock(library):
    job = jobs.create_research_job(QUESTION, frames('甲'))
    def interrupt(stage):
        raise KeyboardInterrupt()
    with pytest.raises(KeyboardInterrupt):
        jobs.resume_research_job(job['job_id'], check_interrupt=interrupt)
    assert jobs.list_research_jobs()[0]['status'] == 'paused'
    assert jobs.resume_research_job(job['job_id'])['status'] == 'completed'


def test_duplicate_click_cannot_enter_running_job(library, monkeypatch):
    job = jobs.create_research_job(QUESTION, frames('甲'))
    nested = []
    def retrieve(*args, **kwargs):
        nested.append(jobs.resume_research_job(job['job_id']))
        return {'answer': '已查', 'status': 'success'}
    monkeypatch.setattr(task, 'retrieve_documents', retrieve)
    assert jobs.resume_research_job(job['job_id'])['status'] == 'completed'
    assert len(nested) == 1 and nested[0]['status'] == 'busy'


def test_public_restore_capability_copies_only_this_job_and_rebinds_all_paths(tmp_path, monkeypatch):
    monkeypatch.setenv('FINANCIAL_DEPLOYMENT', 'public')
    monkeypatch.setenv('FINANCIAL_DATA_DIR', str(tmp_path))
    try:
        configure_public_session('a' * 32)
        source = get_data_root() / 'uploads' / 'source.csv'
        source.parent.mkdir(parents=True)
        source.write_text('financial source', encoding='utf-8')
        data = frames('甲', '乙')
        for frame in data.values():
            frame['metric_provenance'] = [{'revenue': {'file_path': str(source)}}] * 2
        reports = [{'file_path': str(source), 'file_name': source.name, 'company_name': '甲', 'report_year': 2024}]
        def retrieve(question, context, reports, **kwargs):
            return {'answer': '冻结原文', 'status': 'success', 'citations': [{'file_path': reports[0]['file_path'], 'snippet': '已查'}]}
        monkeypatch.setattr(task, 'retrieve_documents', retrieve)
        job = jobs.create_research_job(QUESTION, data, reports=reports)
        hidden = jobs.create_research_job(QUESTION.replace('2024', '2025'), data)
        def pause(stage):
            if stage.startswith('企业检索检查点已保存'):
                raise BudgetExceeded()
        assert jobs.resume_research_job(job['job_id'], check_interrupt=pause)['completed'] == 1
        issued = jobs.issue_recovery_code(job['job_id'])
        code = issued['code']
        assert len(code) == 43
        disk = '\n'.join(p.read_text(encoding='utf-8') for p in tmp_path.rglob('*.json'))
        assert code not in disk
        configure_public_session('b' * 32)
        with pytest.raises(ValueError):
            jobs.resume_research_job(job['job_id'])
        assert jobs.list_research_jobs() == []
        restored = jobs.restore_research_job(code)
        assert restored['completed'] == 1 and restored['restored']
        assert len(jobs.list_research_jobs()) == 1
        assert jobs.restore_research_job(code)['job_id'] == job['job_id']
        assert not (get_data_root() / 'research_jobs' / hidden['job_id']).exists()
        done = jobs.resume_research_job(job['job_id'])
        for raw in jobs._walk_paths(done['result']):
            assert Path(raw).is_relative_to(get_data_root())
            assert resolve_stored_file(raw).is_file()
        assert done['result']['data_snapshot']['甲'][1]['revenue'] == 120
    finally:
        configure_workspace('main')


def test_expiry_and_code_rotation(library, monkeypatch):
    job = jobs.create_research_job(QUESTION, frames('甲'))
    first = jobs.issue_recovery_code(job['job_id'])
    second = jobs.issue_recovery_code(job['job_id'])
    with pytest.raises(ValueError, match='恢复码无效'):
        jobs.restore_research_job(first['code'])
    monkeypatch.setattr(jobs.time, 'time', lambda: second['expires_at'] + 1)
    with pytest.raises(ValueError, match='恢复码无效'):
        jobs.restore_research_job(second['code'])


def test_delivered_is_not_global_denial_of_completed_result(library):
    job = jobs.create_research_job(QUESTION, frames('甲'))
    assert not jobs.mark_research_job_delivered(job['job_id'])
    done = jobs.resume_research_job(job['job_id'])
    assert jobs.mark_research_job_delivered(job['job_id'])
    assert jobs.mark_research_job_delivered(job['job_id'])
    assert jobs.resume_research_job(job['job_id'])['result'] == done['result']


def test_no_credentials_and_no_unowned_original_copy(library, tmp_path):
    foreign = tmp_path / 'other' / 'secret.csv'
    foreign.parent.mkdir()
    foreign.write_text('not authorized')
    data = frames('甲')
    data['甲']['api_key'] = 'private-key-never-store'
    data['甲']['metric_provenance'] = [{'revenue': {'file_path': str(foreign)}}] * 2
    job = jobs.create_research_job(QUESTION, data, reports=[{'file_path': str(foreign), 'password': 'hidden'}])
    raw = (library / 'research_jobs' / job['job_id'] / 'job.json').read_text(encoding='utf-8')
    assert 'private-key' not in raw and 'password' not in raw and 'not authorized' not in raw
    state = json.loads(raw)
    assert state['data']['甲'][0]['metric_provenance']['revenue']['file_path'] == ''
    assert '待复核' in state['data']['甲'][0]['metric_provenance']['revenue']['source_availability']


def test_invalid_identifiers_limits_and_unsupported_requests(library, monkeypatch):
    for identity in ['../other', 'a' * 31, 'Z' * 32]:
        with pytest.raises(ValueError):
            jobs.resume_research_job(identity)
    with pytest.raises(ValueError, match='300家'):
        jobs.create_research_job(QUESTION, frames(*(str(i) for i in range(301))))
    monkeypatch.setattr(jobs, 'MAX_JOBS', 1)
    one = jobs.create_research_job('请讲个故事', frames('甲'))
    assert jobs.resume_research_job(one['job_id'])['result']['status'] == 'unsupported'
    with pytest.raises(ValueError, match='20个'):
        jobs.create_research_job(QUESTION, frames('甲'))


def test_shared_public_workspace_is_read_only(tmp_path, monkeypatch):
    monkeypatch.setenv('FINANCIAL_DEPLOYMENT', 'public')
    monkeypatch.setenv('FINANCIAL_DATA_DIR', str(tmp_path))
    configure_workspace('main')
    with pytest.raises(PermissionError):
        jobs.create_research_job(QUESTION, frames('甲'))
    assert not (tmp_path / 'main' / 'research_jobs').exists()


def test_state_size_counts_toward_workspace_quota(library, monkeypatch):
    monkeypatch.setattr(jobs, 'MAX_WORKSPACE_BYTES', 1000)
    data = frames('甲')
    data['甲']['note'] = 'x' * 1500
    with pytest.raises(ValueError, match='200 MB'):
        jobs.create_research_job(QUESTION, data)
    assert jobs.list_research_jobs() == []
    assert jobs._workspace_bytes() <= 1000


def test_checkpoint_growth_keeps_previous_valid_checkpoint_and_discards_index(library, monkeypatch):
    job = jobs.create_research_job(QUESTION, frames('甲', '乙'))
    ceiling = jobs._workspace_bytes() + 2200
    monkeypatch.setattr(jobs, 'MAX_WORKSPACE_BYTES', ceiling)
    calls = []
    def retrieve(question, context, index_root, **kwargs):
        name = context['companies'][0]
        calls.append(name)
        index_root.mkdir(exist_ok=True)
        (index_root / 'cache.json').write_text('derived cache' * 500)
        return {'answer': name if name == '甲' else 'x' * 5000, 'status': 'success'}
    monkeypatch.setattr(task, 'retrieve_documents', retrieve)
    paused = jobs.resume_research_job(job['job_id'])
    assert paused['status'] == 'paused' and paused['completed'] == 1
    assert jobs._workspace_bytes() <= ceiling
    assert not (library / 'research_jobs' / job['job_id'] / 'index').exists()
    monkeypatch.setattr(task, 'retrieve_documents', lambda *args, **kwargs: {'answer': '较短结果', 'status': 'success'})
    monkeypatch.setattr(jobs, 'MAX_WORKSPACE_BYTES', 200 * 1024 * 1024)
    done = jobs.resume_research_job(job['job_id'])
    assert done['status'] == 'completed' and done['result']['rows'][0]['explanation'] == '甲'






